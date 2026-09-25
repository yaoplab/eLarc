"""Sorties fichier du rapport « Matières par classe » : Excel et sauvegarde JSON.

Excel : une feuille par page de données, logo + titre en tête, police Arial,
prête à imprimer (A4 paysage, ajustée en largeur, en-têtes de colonnes répétés).
Sauvegarde JSON : lignes brutes de la base (identifiants inclus) pour cette
classe et ce trimestre — de quoi reconstituer l'état si une modification tourne mal.
"""
from __future__ import annotations

import json
import os

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from LarcConfig.common.matieres_report import SCHOOL_NAME, Report, Table

LOGO_PATH = os.path.join(os.path.dirname(os.path.dirname(__file__)), "img", "logoAEC.png")

_FONT = "Arial"
_BLUE = "1F4E9C"
_LIGHT = "DCE8F8"
_GREY = "6B7280"
_WARN_BG = "FDE8E8"
_THIN = Side(style="thin", color="C9CED6")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_MAX_COL_WIDTH = 46


def export_xlsx(report: Report, path: str) -> str:
    wb = Workbook()
    wb.remove(wb.active)
    for section in report.sections:
        _write_sheet(wb.create_sheet(section.title[:31]), report, section)
    wb.save(path)
    return path


def _write_sheet(ws, report: Report, section):
    n_cols = max(len(t.columns) for t in section.tables)
    ws.sheet_view.showGridLines = False
    if os.path.exists(LOGO_PATH):
        img = XlImage(LOGO_PATH)
        ratio = 60 / img.height
        img.height, img.width = 60, int(img.width * ratio)
        ws.add_image(img, "A1")
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 20
    # Le titre commence en colonne C pour laisser la place au logo.
    first_text_col = min(3, n_cols)
    ws.cell(row=1, column=first_text_col, value=report.title).font = Font(name=_FONT, size=16, bold=True, color=_BLUE)
    ws.cell(row=2, column=first_text_col, value=f"{report.subtitle}  ·  {section.title}").font = Font(
        name=_FONT, size=10, color=_GREY)
    ws.cell(row=3, column=first_text_col,
            value=f"{SCHOOL_NAME} — édité le {report.generated:%d/%m/%Y à %H:%M}").font = Font(
        name=_FONT, size=9, italic=True, color=_GREY)

    widths = [8] * n_cols
    row = 5
    header_rows = []
    for table in section.tables:
        row = _write_table(ws, table, row, widths, header_rows) + 1
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 10), _MAX_COL_WIDTH)
    if len(section.tables) == 1:   # figer / répéter l'en-tête seulement s'il n'y a qu'un tableau
        ws.freeze_panes = ws.cell(row=header_rows[0] + 1, column=1)
        ws.print_title_rows = f"{header_rows[0]}:{header_rows[0]}"

    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.oddHeader.left.text = f"&\"{_FONT},Bold\"{report.title}"
    ws.oddHeader.right.text = f"&\"{_FONT},Regular\"{report.subtitle}"
    ws.oddFooter.left.text = f"&\"{_FONT},Regular\"{SCHOOL_NAME}"
    ws.oddFooter.right.text = f"&\"{_FONT},Regular\"Page &P / &N"


def _write_table(ws, table: Table, row: int, widths: list[int], header_rows: list[int]) -> int:
    n = len(table.columns)
    ws.cell(row=row, column=1, value=table.title).font = Font(name=_FONT, size=12, bold=True, color=_BLUE)
    ws.row_dimensions[row].height = 20
    row += 1
    header_rows.append(row)
    for col, text in enumerate(table.columns, 1):
        c = ws.cell(row=row, column=col, value=text)
        c.font = Font(name=_FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=_BLUE)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER
        widths[col - 1] = max(widths[col - 1], len(text))
    ws.row_dimensions[row].height = 24
    row += 1

    if not table.rows:
        ws.cell(row=row, column=1, value=table.empty_text).font = Font(name=_FONT, size=10, italic=True, color=_GREY)
        return row + 1
    for i, values in enumerate(table.rows):
        flag = table.flags[i] if i < len(table.flags) else ""
        for col, v in enumerate(values, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.font = Font(name=_FONT, size=10, color="B42318" if flag == "warn" and col == n else "000000")
            c.border = _BORDER
            align = table.aligns[col - 1] if col - 1 < len(table.aligns) else "l"
            c.alignment = Alignment(horizontal="center" if align == "c" else "left",
                                    vertical="center", wrap_text=True)
            if flag == "warn":
                c.fill = PatternFill("solid", fgColor=_WARN_BG)
            elif i % 2:
                c.fill = PatternFill("solid", fgColor="F5F8FD")
            widths[col - 1] = max(widths[col - 1], min(len(str(v)), _MAX_COL_WIDTH))
        row += 1
    if table.note:
        ws.cell(row=row, column=1, value=table.note).font = Font(name=_FONT, size=9, italic=True, color=_GREY)
        row += 1
    return row


def export_backup_json(report: Report, path: str) -> str:
    with open(path, "w", encoding="utf-8") as f:
        json.dump(report.backup, f, ensure_ascii=False, indent=2, default=str)
    return path
