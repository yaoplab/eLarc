"""Export / import Excel des types d'événements (larcauth_event_type_config).

Export : un classeur avec un mode d'emploi et une feuille par langue (Français, English),
mis en page (logo, titre, niveaux, colonnes verrouillées sauf Libellé et Actif).
Import : relit ce classeur et prépare un PLAN de modifications, vérifié règle par règle,
avant toute écriture. Principe gabarit : on ne crée et on ne supprime jamais de ligne ;
on ne fait que renommer ou activer/désactiver des lignes existantes.
"""
from __future__ import annotations

import datetime
import os
from dataclasses import dataclass, field

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

from LarcConfig.common.db_access import (
    FREE_LABEL_PREFIX, get_event_type_tree, set_event_type_active, set_event_type_label,
)

SHEETS = {2: "Français", 1: "English"}
HEADERS = ["ID", "Niveau", "Libellé", "Utilisé", "Actif", "Emplacement"]
COL_ID, COL_LEVEL, COL_LABEL, COL_USED, COL_ACTIVE, COL_SLOT = range(1, 7)
FIRST_DATA_ROW = 6
HELP_SHEET = "Mode d'emploi"
_LOGO = os.path.join(os.path.dirname(os.path.dirname(__file__)), "img", "logoAEC.png")

_BLUE = "1F4E9C"
_LIGHT = "DCE8F8"
_GREY = "8A8F98"
_THIN = Side(style="thin", color="C9CED6")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


# --------------------------------------------------------------------------- export
def export_types_xlsx(path: str, trees: dict | None = None, author: str = "") -> str:
    """Écrit le classeur. `trees` = {fk_language: lignes de get_event_type_tree(include_free=True)}."""
    if trees is None:
        trees = {lang: get_event_type_tree(lang, include_free=True) for lang in SHEETS}
    wb = Workbook()
    _write_help(wb.active, author)
    for lang, name in SHEETS.items():
        _write_sheet(wb.create_sheet(name), name, trees.get(lang, []))
    wb.save(path)
    return path


def _add_logo(ws, height: int = 62):
    if os.path.exists(_LOGO):
        img = XlImage(_LOGO)
        ratio = height / img.height
        img.height, img.width = height, int(img.width * ratio)
        ws.add_image(img, "A1")


def _write_sheet(ws, name: str, rows: list):
    _add_logo(ws)
    ws.merge_cells("C1:F1")
    ws["C1"] = f"Types d'événements — {name}"
    ws["C1"].font = Font(name="Calibri", size=18, bold=True, color=_BLUE)
    ws["C1"].alignment = Alignment(vertical="center")
    ws.merge_cells("C2:F2")
    ws["C2"] = (f"Exporté le {datetime.date.today():%d/%m/%Y} — modifiez uniquement les colonnes "
                "« Libellé » et « Actif », puis importez le fichier depuis LarcConfig.")
    ws["C2"].font = Font(size=10, italic=True, color=_GREY)
    ws.row_dimensions[1].height = 48
    ws.row_dimensions[2].height = 18

    header_row = FIRST_DATA_ROW - 1
    for col, text in enumerate(HEADERS, 1):
        c = ws.cell(row=header_row, column=col, value=text)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=_BLUE)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = _BORDER
    ws.row_dimensions[header_row].height = 22

    for i, r in enumerate(rows):
        row = FIRST_DATA_ROW + i
        free = bool(r.get('is_free'))
        level = r['depth'] + 1
        values = [r['id'], f"N{level}", r['label'] or '', r.get('usage', 0) or None,
                  "Oui" if r['enabled'] else "Non", "Libre" if free else ""]
        for col, v in enumerate(values, 1):
            c = ws.cell(row=row, column=col, value=v)
            c.border = _BORDER
            c.protection = Protection(locked=col not in (COL_LABEL, COL_ACTIVE))
            if col in (COL_ID, COL_LEVEL, COL_USED, COL_ACTIVE, COL_SLOT):
                c.alignment = Alignment(horizontal="center")
            if col == COL_LABEL:
                c.alignment = Alignment(indent=r['depth'] * 2)
            if level == 1 and r['enabled']:
                c.font = Font(bold=True, color=_BLUE)
                c.fill = PatternFill("solid", fgColor=_LIGHT)
            elif free:
                c.font = Font(italic=True, color=_GREY)
            elif not r['enabled']:
                c.font = Font(color=_GREY)
            if r.get('usage') and col == COL_USED:
                c.font = Font(bold=True, color="B45309")

    last = FIRST_DATA_ROW + max(len(rows), 1) - 1
    dv = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=False,
                        errorTitle="Valeur invalide", error="Choisissez Oui ou Non.")
    ws.add_data_validation(dv)
    dv.add(f"E{FIRST_DATA_ROW}:E{last}")

    for col, width in zip("ABCDEF", (10, 9, 58, 10, 9, 14)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = ws.cell(row=FIRST_DATA_ROW, column=1)
    ws.auto_filter.ref = f"A{header_row}:F{last}"
    ws.protection.sheet = True                 # ID, Niveau, Utilisé, Emplacement verrouillés
    ws.protection.autoFilter = False           # …mais filtre et tri restent possibles
    ws.protection.sort = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{header_row}:{header_row}"
    ws.oddFooter.center.text = "Page &P / &N"
    ws.oddHeader.right.text = f"Types d'événements — {name}"


def _write_help(ws, author: str):
    ws.title = HELP_SHEET
    _add_logo(ws)
    ws.merge_cells("C1:H1")
    ws["C1"] = "Types d'événements — mode d'emploi"
    ws["C1"].font = Font(size=18, bold=True, color=_BLUE)
    ws.row_dimensions[1].height = 48
    lines = [
        ("Ce classeur contient tous les types d'événements (absences, retards, événements, sorties).", False),
        ("Une feuille par langue : Français et English.", False),
        ("", False),
        ("Ce que vous pouvez modifier", True),
        ("• Colonne « Libellé » : le nom du type (impossible si le type est déjà utilisé).", False),
        ("• Colonne « Actif » : Oui / Non (impossible de désactiver un type déjà utilisé).", False),
        ("", False),
        ("Ce que vous ne pouvez pas modifier", True),
        ("• ID, Niveau, Utilisé, Emplacement : colonnes verrouillées.", False),
        ("• On ne crée ni ne supprime jamais de ligne : pour un NOUVEAU type, prenez une ligne « Libre »,", False),
        (f"  donnez-lui un nom (à la place de « {FREE_LABEL_PREFIX}… ») et mettez Actif = Oui.", False),
        ("", False),
        ("Comment importer", True),
        ("1. Enregistrez ce fichier (format .xlsx).", False),
        ("2. Dans LarcConfig > Types d'événements, cliquez sur « Importer Excel ».", False),
        ("3. Vérifiez l'aperçu des modifications (et des refus éventuels), puis confirmez.", False),
        ("", False),
        ("Règles vérifiées à l'import", True),
        ("• Un type utilisé par des événements ne peut être ni renommé ni désactivé.", False),
        ("• Un type ne peut être actif que si son parent l'est ; un parent ne peut être désactivé", False),
        ("  tant qu'un de ses sous-types est actif.", False),
        ("• Un emplacement libre doit être renommé avant d'être activé.", False),
    ]
    for i, (text, title) in enumerate(lines, 3):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(bold=True, size=12, color=_BLUE) if title else Font(size=11)
    if author:
        note = ws.cell(row=len(lines) + 4, column=1, value=f"Export généré par {author}.")
        note.font = Font(italic=True, color=_GREY)
    ws.column_dimensions["A"].width = 14      # le texte déborde sur les cellules vides voisines


# --------------------------------------------------------------------------- import
@dataclass
class Change:
    lang: int
    id: int
    kind: str                # 'label' | 'activate' | 'deactivate'
    old: str
    new: str


@dataclass
class ImportPlan:
    changes: list = field(default_factory=list)
    rejected: list = field(default_factory=list)   # (lang, id, raison)

    def summary(self) -> str:
        return f"{len(self.changes)} modification(s) prévue(s), {len(self.rejected)} ligne(s) refusée(s)."


def read_types_xlsx(path: str) -> dict:
    """Lit les feuilles Français / English : [{'id', 'label', 'active', 'line'}] par langue."""
    wb = load_workbook(path, data_only=True)
    out: dict = {}
    for lang, name in SHEETS.items():
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        header_row = next(
            (r for r in range(1, 15) if str(ws.cell(row=r, column=1).value).strip() == "ID"), None)
        if header_row is None:
            raise ValueError(f"Feuille « {name} » : en-tête « ID » introuvable.")
        rows = []
        for r in range(header_row + 1, ws.max_row + 1):
            raw_id = ws.cell(row=r, column=COL_ID).value
            if raw_id in (None, ""):
                continue
            rows.append({'id': raw_id, 'label': ws.cell(row=r, column=COL_LABEL).value,
                         'active': ws.cell(row=r, column=COL_ACTIVE).value, 'line': r})
        out[lang] = rows
    if not out:
        raise ValueError("Aucune feuille « Français » ou « English » dans ce fichier.")
    return out


def _parse_active(value):
    if isinstance(value, bool):
        return value
    text = str(value if value is not None else "").strip().lower()
    if text in ("oui", "o", "1", "true", "vrai", "x", "yes"):
        return True
    if text in ("non", "n", "0", "false", "faux", "no"):
        return False
    return None


def plan_import(sheets: dict, current: dict) -> ImportPlan:
    """Compare le fichier à l'état actuel et applique les règles. N'écrit rien."""
    plan = ImportPlan()
    for lang, file_rows in sheets.items():
        cur_by_id = {r['id']: r for r in current.get(lang, [])}
        desired: dict = {}
        seen: set = set()
        for fr in file_rows:
            try:
                tid = int(fr['id'])
            except (TypeError, ValueError):
                plan.rejected.append((lang, 0, f"ligne {fr['line']} : ID « {fr['id']} » invalide"))
                continue
            row = cur_by_id.get(tid)
            if row is None:
                plan.rejected.append((lang, tid, "ID inconnu (on ne crée jamais de ligne)"))
                continue
            if tid in seen:
                plan.rejected.append((lang, tid, "ID présent deux fois dans la feuille"))
                continue
            seen.add(tid)
            label = str(fr['label'] if fr['label'] is not None else "").strip()
            active = _parse_active(fr['active'])
            if not label:
                plan.rejected.append((lang, tid, "libellé vide"))
                continue
            if active is None:
                plan.rejected.append(
                    (lang, tid, f"« Actif » doit valoir Oui ou Non (reçu : {fr['active']!r})"))
                continue
            used = row.get('usage', 0)
            old_label = row['label'] or ""
            if label != old_label and used:
                plan.rejected.append((lang, tid, f"utilisé par {used} événement(s) : renommage interdit"))
                label = old_label
            if not active and row['enabled'] and used:
                plan.rejected.append(
                    (lang, tid, f"utilisé par {used} événement(s) : désactivation interdite"))
                active = True
            if active and not row['enabled'] and label.startswith(FREE_LABEL_PREFIX):
                plan.rejected.append((lang, tid, "emplacement libre : renommez-le avant de l'activer"))
                active = False
            desired[tid] = (label, active)
        _enforce_hierarchy(lang, cur_by_id, desired, plan)
        for tid in sorted(desired):
            row, (label, active) = cur_by_id[tid], desired[tid]
            if label != (row['label'] or ""):
                plan.changes.append(Change(lang, tid, 'label', row['label'] or "", label))
            if active != row['enabled']:
                plan.changes.append(Change(lang, tid, 'activate' if active else 'deactivate',
                                           "Non" if active else "Oui", "Oui" if active else "Non"))
    return plan


def _enforce_hierarchy(lang, cur_by_id, desired, plan):
    """Un type actif exige un parent actif ; un parent ne se désactive pas s'il a un fils actif."""
    def state(tid):
        if tid in desired:
            return desired[tid][1]
        return cur_by_id[tid]['enabled'] if tid in cur_by_id else True

    for _pass in range(10):                               # point fixe (profondeur <= 4)
        changed = False
        for tid in sorted(desired):
            label, active = desired[tid]
            row = cur_by_id[tid]
            parent = row.get('parent_id')
            if active and not row['enabled'] and parent is not None and not state(parent):
                desired[tid] = (label, False)
                plan.rejected.append((lang, tid, "parent inactif : activez d'abord le parent"))
                changed = True
                continue
            if not active and row['enabled']:
                kids = [k for k, r in cur_by_id.items() if r.get('parent_id') == tid and state(k)]
                if kids:
                    desired[tid] = (label, True)
                    plan.rejected.append(
                        (lang, tid, "un sous-type est encore actif : désactivez-le d'abord"))
                    changed = True
        if not changed:
            break


def apply_plan(plan: ImportPlan):
    """Applique le plan via db_access (garde-fous compris). Retourne (nb appliquées, échecs)."""
    done, failed = 0, []
    # renommages et activations : parents d'abord ; désactivations : enfants d'abord
    order = sorted(plan.changes,
                   key=lambda c: (c.kind == 'deactivate', -c.id if c.kind == 'deactivate' else c.id))
    for ch in order:
        if ch.kind == 'label':
            ok = set_event_type_label(ch.id, ch.new)
        else:
            ok = set_event_type_active(ch.id, ch.kind == 'activate')
        if ok:
            done += 1
        else:
            failed.append(f"{ch.id} ({SHEETS[ch.lang]}) : {ch.kind} refusé")
    return done, failed
