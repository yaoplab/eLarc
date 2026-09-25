"""Tests du rapport « Matières par classe » (impression, Excel, JSON) — sans base."""
from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from LarcConfig.common import matieres_export, matieres_report as mr


def _subject(id_, label, group, enabled=True, niv_sup=False, cross=False, enrolled=0, tl="LABOE", tf="Anselme"):
    return {'id': id_, 'label': label, 'enabled': enabled, 'niv_sup': niv_sup, 'cross_track': cross,
            'group_label': group, 'nr_group_in_pgm': 1, 'teacher_last_name': tl,
            'teacher_first_name': tf, 'enrolled_count': enrolled}


def _patch(monkeypatch, subjects, groups, students, enrolments, slots=(), links=()):
    monkeypatch.setattr(mr.db_enrolment, 'get_classroom_subjects', lambda c, t: list(subjects))
    monkeypatch.setattr(mr.db_enrolment, 'get_subject_groups', lambda c, t: list(groups))
    monkeypatch.setattr(mr.db_enrolment, 'get_active_students', lambda c: list(students))
    monkeypatch.setattr(mr.db_enrolment, 'get_student_enrolments', lambda c, t: list(enrolments))
    monkeypatch.setattr(mr.db_othersubjects, 'get_classroom_othersubjects', lambda c, t: list(slots))
    monkeypatch.setattr(mr.db_othersubjects, 'get_othersubject_links', lambda c, t: list(links))
    monkeypatch.setattr(mr.db_othersubjects, 'get_supervisors',
                        lambda: [{'id': 7, 'last_name': 'DUPUY', 'first_name': 'Thierry'}])


def _enrol(sid, cts, label, group, niv_sup=False):
    return {'fk_student_id': sid, 'cts_id': cts, 'label': label, 'niv_sup': niv_sup,
            'cross_track': False, 'couleur': None, 'nr_group_in_pgm': group}


STUDENTS = [{'aecuser_ptr_id': 1, 'last_name': 'ALAOFE', 'first_name': 'Naofal'}]
GROUPS = [{'nr_group_in_pgm': 1, 'group_label': 'Langue'}]
PEI = {'id': 5, 'label': 'PEI-3', 'program_sigle': 'PEI'}
DP = {'id': 5, 'label': 'DP-2Fr', 'program_sigle': 'DPFr'}


def test_only_active_subjects_are_listed_and_counted(monkeypatch):
    subs = [_subject(1, 'Français', 'Langue'), _subject(2, 'Matiere_supl 1', 'Langue', enabled=False)]
    _patch(monkeypatch, subs, GROUPS, STUDENTS, [])
    t = mr.build_report(PEI, 1, 1, '2026-2027').sections[0].tables[0]
    assert [r[1] for r in t.rows] == ['Français']
    assert '1 emplacement(s) inactif(s)' in t.note
    # ni Niveau (hors DP) ni Piste croisée (jamais utilisée)
    assert t.columns == ['Groupe', 'Matière', 'Enseignant', 'Inscrits']


def test_dp_has_level_column_and_cross_track_only_when_used(monkeypatch):
    subs = [_subject(1, 'Anglais NS', 'Langues', niv_sup=True, cross=True)]
    _patch(monkeypatch, subs, GROUPS, STUDENTS, [])
    t = mr.build_report(DP, 1, 1).sections[0].tables[0]
    assert t.columns == ['Groupe', 'Matière', 'Niveau', 'Enseignant', 'Inscrits', 'Piste croisée']
    assert t.rows[0][2] == 'NS' and t.rows[0][-1] == 'Oui'


def test_student_cells_and_dp_status(monkeypatch):
    # un seul élève avec 2 matières : non conforme en DP
    enr = [_enrol(1, 1, 'Français', 1, niv_sup=True), _enrol(1, 2, 'Maths NM', 1, niv_sup=False)]
    _patch(monkeypatch, [], GROUPS, STUDENTS, enr)
    t = mr.build_report(DP, 1, 1).sections[1].tables[0]
    assert t.columns == ['Élève', 'Langue', 'Statut DP']
    assert t.rows[0][1] == 'Français\nMaths NM'   # une matière par ligne, pas de « + »
    assert t.rows[0][2].startswith('Non conforme — ')
    assert t.flags == ['warn']


def test_no_ns_tag_added_but_cross_track_marked(monkeypatch):
    enr = [_enrol(1, 1, 'Anglais NS', 1, niv_sup=True), dict(_enrol(1, 2, 'Sciences', 1), cross_track=True)]
    _patch(monkeypatch, [], GROUPS, STUDENTS, enr)
    t = mr.build_report(DP, 1, 1).sections[1].tables[0]
    assert t.rows[0][1] == 'Anglais NS\nSciences (*)'


def test_no_dp_status_column_outside_dp(monkeypatch):
    _patch(monkeypatch, [], GROUPS, STUDENTS, [_enrol(1, 1, 'Français', 1)])
    assert mr.build_report(PEI, 1, 1).sections[1].tables[0].columns == ['Élève', 'Langue']


def test_other_subjects_grid_shows_tutor_only_when_enrolled(monkeypatch):
    slots = [{'id': 10, 'label': 'CAS', 'enabled': True, 'supervisor_last_name': 'DUPUY',
              'supervisor_first_name': 'Thierry', 'enrolled_count': 1},
             {'id': 11, 'label': 'Inactif', 'enabled': False, 'supervisor_last_name': None,
              'supervisor_first_name': None, 'enrolled_count': 0}]
    links = [{'fk_student_id': 1, 'fk_termothersubject_id': 10, 'enabled': True, 'ref_teacher': 7}]
    _patch(monkeypatch, [], GROUPS, STUDENTS, [], slots, links)
    slots_t, grid = mr.build_report(DP, 1, 1).sections[2].tables
    assert slots_t.rows == [['CAS', 'DUPUY Thierry', '1']]
    assert grid.columns == ['Élève', 'CAS'] and grid.rows == [['ALAOFE Naofal', 'DUPUY Thierry']]


def test_other_subjects_without_active_slot_has_single_empty_table(monkeypatch):
    _patch(monkeypatch, [], GROUPS, STUDENTS, [], slots=[])
    tables = mr.build_report(PEI, 1, 1).sections[2].tables
    assert len(tables) == 1 and tables[0].rows == []


@pytest.fixture
def report(monkeypatch):
    subs = [_subject(1, 'Français', 'Langue', enrolled=1)]
    _patch(monkeypatch, subs, GROUPS, STUDENTS, [_enrol(1, 1, 'Français', 1)])
    return mr.build_report(PEI, 1, 1, '2026-2027')


def test_titles_and_file_name(report):
    assert report.title == 'Matières de classe — PEI-3'
    assert report.subtitle == 'PEI · Trimestre 1 · 2026-2027'
    assert report.file_stem.startswith('Matieres_PEI-3_T1_')


def test_excel_is_arial_with_logo_and_three_sheets(report, tmp_path):
    wb = load_workbook(matieres_export.export_xlsx(report, str(tmp_path / "r.xlsx")))
    assert wb.sheetnames == ['Classes & matières', 'Élèves', 'Autres-matières']
    for ws in wb:
        assert len(ws._images) == 1
        assert {c.font.name for row in ws.iter_rows() for c in row if c.value is not None} == {'Arial'}
        assert ws.page_setup.orientation == 'landscape'


def test_backup_json_holds_raw_rows_and_context(report, tmp_path):
    path = matieres_export.export_backup_json(report, str(tmp_path / "b.json"))
    with open(path, encoding="utf-8") as f:
        data = json.load(f)
    assert data['format'] == mr.BACKUP_FORMAT
    assert data['classroom']['label'] == 'PEI-3' and data['term'] == {'number': 1, 'id': 1}
    assert data['classroom_subjects'][0]['id'] == 1


def test_pdf_is_written(report, tmp_path):
    from PySide6.QtWidgets import QApplication
    from LarcConfig.common import matieres_print
    QApplication.instance() or QApplication([])
    path = str(tmp_path / "r.pdf")
    assert matieres_print.save_pdf(report, path) is True
    with open(path, "rb") as f:
        assert f.read(5) == b"%PDF-"


def test_html_uses_arial_and_escapes_text(report):
    from LarcConfig.common import matieres_print
    report.sections[0].tables[0].rows[0][1] = "<b>Fr</b> & co"
    html = matieres_print.report_html(report)
    assert 'font-family:\'Arial\'' in html
    assert "&lt;b&gt;Fr&lt;/b&gt; &amp; co" in html


def test_student_group_columns_are_centered(monkeypatch):
    _patch(monkeypatch, [], GROUPS, STUDENTS, [_enrol(1, 1, 'Mathématiques NM - Analyse et Approches', 1)])
    t = mr.build_report(DP, 1, 1).sections[1].tables[0]
    assert t.rows[0][1] == 'Mathématiques NM - Analyse et Approches'   # libellé complet
    assert t.aligns == ['l', 'c', 'l']


def test_html_breaks_lines_inside_cells(report):
    from LarcConfig.common import matieres_print
    report.sections[0].tables[0].rows[0][1] = "A" + chr(10) + "B"
    assert "A<br>B" in matieres_print.report_html(report)
