"""Tests export / import Excel des types d'événements (sans base : lignes factices)."""
from __future__ import annotations

import copy

from openpyxl import load_workbook

from LarcConfig.common import types_excel as tx

FR = 2


def _rows():
    def r(id_, label, depth, parent, enabled=True, usage=0, free=False):
        return {'id': id_, 'label': label, 'depth': depth, 'parent_id': parent,
                'enabled': enabled, 'usage': usage, 'is_free': free, 'code': f"c{id_}"}
    return [
        r(21000, "Absent de l'école", 0, None, usage=2),
        r(21100, "Justifiée", 1, 21000, usage=2),
        r(21200, "Injustifiée", 1, 21000),
        r(21300, "Ancien motif", 1, 21000, enabled=False),
        r(21400, "Categorie_Niv21000_Level_4", 1, 21000, enabled=False, free=True),
        r(21410, "Categorie_Niv21400_Level_1", 2, 21400, enabled=False, free=True),
    ]


def _sheet(rows, overrides=None):
    """Feuille telle que relue depuis Excel, à partir de lignes courantes (+ modifications)."""
    overrides = overrides or {}
    out = []
    for i, r in enumerate(rows):
        row = {'id': r['id'], 'label': r['label'], 'active': "Oui" if r['enabled'] else "Non",
               'line': 6 + i}
        row.update(overrides.get(r['id'], {}))
        out.append(row)
    return out


def _plan(overrides, rows=None):
    rows = rows or _rows()
    return tx.plan_import({FR: _sheet(rows, overrides)}, {FR: copy.deepcopy(rows)})


def test_export_then_read_gives_no_change(tmp_path):
    path = str(tmp_path / "types.xlsx")
    tx.export_types_xlsx(path, {FR: _rows(), 1: []}, author="Test")
    sheets = tx.read_types_xlsx(path)
    plan = tx.plan_import(sheets, {FR: _rows(), 1: []})
    assert plan.changes == [] and plan.rejected == []


def test_export_layout_logo_title_protection_and_validation(tmp_path):
    path = str(tmp_path / "types.xlsx")
    tx.export_types_xlsx(path, {FR: _rows(), 1: []})
    wb = load_workbook(path)
    assert wb.sheetnames == [tx.HELP_SHEET, "Français", "English"]
    ws = wb["Français"]
    assert ws["C1"].value == "Types d'événements — Français"
    assert ws._images, "le logo doit être présent"
    assert [ws.cell(row=5, column=c).value for c in range(1, 7)] == tx.HEADERS
    assert ws.protection.sheet is True
    assert ws.cell(row=6, column=tx.COL_ID).protection.locked is True
    assert ws.cell(row=6, column=tx.COL_LABEL).protection.locked is False
    assert ws.cell(row=6, column=tx.COL_ACTIVE).protection.locked is False
    assert ws.freeze_panes == "A6"
    assert ws.cell(row=9, column=tx.COL_ACTIVE).value == "Non"        # 4e type : désactivé
    assert ws.cell(row=10, column=tx.COL_SLOT).value == "Libre"       # emplacement libre
    assert ws.data_validations.dataValidation[0].formula1 == '"Oui,Non"'


def test_rename_and_toggle_are_planned(tmp_path):
    plan = _plan({21200: {'label': "Injustifiée (sans motif)"}, 21300: {'active': "Oui"}})
    kinds = {(c.id, c.kind) for c in plan.changes}
    assert kinds == {(21200, 'label'), (21300, 'activate')}
    assert plan.rejected == []


def test_used_type_cannot_be_renamed_or_disabled():
    plan = _plan({21100: {'label': "Autre nom", 'active': "Non"}})
    reasons = [r[2] for r in plan.rejected]
    assert any("renommage interdit" in x for x in reasons)
    assert any("désactivation interdite" in x for x in reasons)
    assert not [c for c in plan.changes if c.id == 21100]


def test_free_slot_must_be_renamed_before_activation():
    plan = _plan({21400: {'active': "Oui"}})
    assert any("renommez-le" in r[2] for r in plan.rejected)
    assert plan.changes == []


def test_renamed_free_slot_can_be_activated():
    plan = _plan({21400: {'label': "Nouveau motif", 'active': "Oui"}})
    assert {(c.id, c.kind) for c in plan.changes} == {(21400, 'label'), (21400, 'activate')}
    assert plan.rejected == []


def test_child_cannot_be_active_under_an_inactive_parent():
    # 21410 renommé + activé alors que son parent 21400 (libre) reste inactif
    plan = _plan({21410: {'label': "Sous-motif", 'active': "Oui"}})
    assert any("parent inactif" in r[2] for r in plan.rejected)
    assert not [c for c in plan.changes if c.kind == 'activate']


def test_parent_and_child_activated_together_are_accepted():
    plan = _plan({21400: {'label': "Motif", 'active': "Oui"},
                  21410: {'label': "Sous-motif", 'active': "Oui"}})
    assert plan.rejected == []
    assert {(c.id, c.kind) for c in plan.changes if c.kind == 'activate'} == {(21400, 'activate'),
                                                                                (21410, 'activate')}


def test_unknown_id_empty_label_and_bad_active_are_rejected():
    rows = _rows()
    sheet = _sheet(rows)
    sheet.append({'id': 99999, 'label': "Nouveau", 'active': "Oui", 'line': 20})
    sheet[2]['label'] = "   "
    sheet[1]['active'] = "peut-être"
    plan = tx.plan_import({FR: sheet}, {FR: rows})
    reasons = " | ".join(r[2] for r in plan.rejected)
    assert "ID inconnu" in reasons and "libellé vide" in reasons and "Oui ou Non" in reasons
    assert plan.changes == []


def test_apply_plan_uses_guarded_db_functions(monkeypatch):
    calls = []
    monkeypatch.setattr(tx, "set_event_type_label", lambda i, v: calls.append(("label", i, v)) or True)
    monkeypatch.setattr(tx, "set_event_type_active", lambda i, e: calls.append(("active", i, e)) or True)
    plan = _plan({21400: {'label': "Motif", 'active': "Oui"}, 21200: {'active': "Non"}})
    done, failed = tx.apply_plan(plan)
    assert failed == [] and done == 3
    assert calls[0] == ("label", 21400, "Motif")         # renommage puis activation
    assert ("active", 21400, True) in calls and ("active", 21200, False) in calls
    assert calls.index(("active", 21200, False)) == len(calls) - 1     # désactivations en dernier


def test_apply_plan_reports_refusals(monkeypatch):
    monkeypatch.setattr(tx, "set_event_type_label", lambda i, v: False)
    plan = _plan({21200: {'label': "Autre"}})
    done, failed = tx.apply_plan(plan)
    assert done == 0 and len(failed) == 1
